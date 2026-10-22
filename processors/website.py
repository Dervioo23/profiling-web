"""
Processor untuk data Inbound Website
"""
import io
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


WARNA_HDR = {
    1: "C6EFCE", 2: "DAFDE0", 3: "D9D9D9", 4: "BDD7EE",
    5: "FCE4D6", 6: "FCE4D6", 7: "E2EFDA", 8: "E2EFDA",
    9: "E8D5F5", 10: "E8D5F5", 11: "FFF2CC", 12: "DDEBF7", 13: "DDEBF7",
}
WARNA_DATA = {
    1: "EBF5EB", 2: "F0FAF0", 3: "F2F2F2", 4: "DEEAF1",
    5: "FEF3EE", 6: "FEF3EE", 7: "F2F8EE", 8: "F2F8EE",
    9: "F5EEFA", 10: "F5EEFA", 11: "FFFAE6", 12: "EEF5FB", 13: "EEF5FB",
}

HTML_WARNA_HDR  = {k: f"#{v}" for k, v in WARNA_HDR.items()}
HTML_WARNA_DATA = {k: f"#{v}" for k, v in WARNA_DATA.items()}

# Label kuartal pada sub-header (sesuai notebook asli)
KUARTAL = "Q3"


def _parse_header(df_raw):
    hdr  = df_raw.iloc[1:4].copy()
    data = df_raw.iloc[4:].copy()
    hdr  = hdr.apply(lambda col: col.map(
        lambda x: " ".join(str(x).split()) if pd.notna(x) else np.nan))
    hdr.iloc[0] = hdr.iloc[0].ffill()
    blok  = hdr.iloc[0]
    tahap = hdr.iloc[1].copy()
    for b in blok.dropna().unique():
        idx = blok[blok == b].index
        tahap.loc[idx] = tahap.loc[idx].ffill()
    hdr.iloc[1] = tahap

    def buat_nama(col):
        bagian = [str(x) for x in hdr[col] if pd.notna(x)]
        return "_".join(bagian)

    data.columns = [buat_nama(c) for c in hdr.columns]
    data.columns = [c.replace(" ", "_").replace(".", "") for c in data.columns]
    return data


def _clean_data(data):
    data = data[data.iloc[:, 1].notna()].reset_index(drop=True)
    kolom_prof = [c for c in data.columns if c.startswith("PROFILLING")]
    kolom_off  = [c for c in data.columns if c.startswith("OFFLINE")]
    data_clean = data.drop(columns=kolom_prof + kolom_off)
    kolom_md   = [c for c in data_clean.columns if c.startswith("MD")]
    md = data_clean[kolom_md].apply(pd.to_numeric, errors="coerce")
    mask_nol   = (md.fillna(0) == 0).all(axis=1)
    data_clean = data_clean.loc[~mask_nol].reset_index(drop=True)
    for col in kolom_md:
        data_clean[col] = pd.to_numeric(data_clean[col], errors="coerce").fillna(0)
    return data_clean, kolom_md


def _aggregate(data_clean, kolom_md):
    agg_cols = {}
    for col in ["MD_LEAD", "MD_WL_D", "MD_WL_FLP_Blm_Input_Prospek", "MD_LAPSOPIR"]:
        if col in kolom_md:
            agg_cols[col] = "sum"
    profiling = (
        data_clean.groupby("Nama_Dealer").agg(agg_cols)
        .reset_index()
        .sort_values("MD_LEAD", ascending=False)
        .reset_index(drop=True)
    )
    return profiling


def _build_excel(profiling, bulan, tahun):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inbound Website"

    ws.merge_cells("A1:M1")
    ws["A1"] = "DATA WEBSITE - INBOUND"
    for merge in ["A2:A3", "B2:B3", "C2:C3", "D2:D3",
                  "E2:F2", "G2:H2", "I2:J2", "K2:K3", "L2:M2"]:
        ws.merge_cells(merge)

    ws["A2"] = "Bulan";         ws["B2"] = "Tahun"
    ws["C2"] = "Nama Dealer";   ws["D2"] = "Total Leads"
    ws["E2"] = "KACAB";         ws["G2"] = "FLP"
    ws["I2"] = "Total Workload"; ws["K2"] = "Laporan Sopir"
    ws["L2"] = "Convertion Rate"
    ws["E3"] = "Belum Difollow Up"; ws["F3"] = "Persentase"
    ws["G3"] = "Belum Difollow Up"; ws["H3"] = "Persentase"
    ws["I3"] = f"Target {KUARTAL}";  ws["J3"] = f"Actual {KUARTAL}"
    ws["L3"] = f"Target {KUARTAL}";  ws["M3"] = f"Actual {KUARTAL}"

    for i, row in profiling.iterrows():
        r = i + 4
        lead    = int(row.get("MD_LEAD", 0))
        wl_d    = int(row.get("MD_WL_D", 0))
        wl_flp  = int(row.get("MD_WL_FLP_Blm_Input_Prospek", 0))
        lapsopir= int(row.get("MD_LAPSOPIR", 0))
        ws.cell(r, 1, bulan);  ws.cell(r, 2, tahun)
        ws.cell(r, 3, row["Nama_Dealer"])
        ws.cell(r, 4, lead)
        ws.cell(r, 5, wl_d)
        ws.cell(r, 6, f"=E{r}/D{r}" if lead > 0 else 0)
        ws.cell(r, 7, wl_flp)
        ws.cell(r, 8, f"=G{r}/D{r}" if lead > 0 else 0)
        ws.cell(r, 9, "Maksimal 5%")
        ws.cell(r, 10, f"=F{r}+H{r}")
        ws.cell(r, 11, lapsopir)
        ws.cell(r, 12, "Minimum 2%")
        ws.cell(r, 13, f"=K{r}/D{r}" if lead > 0 else 0)

    thin   = Side(style="thin",   color="000000")
    medium = Side(style="medium", color="000000")
    total_col = 13
    total_row = ws.max_row

    for r in range(4, total_row + 1):
        ws.cell(r, 6).number_format  = "0.00%"
        ws.cell(r, 8).number_format  = "0.00%"
        ws.cell(r, 10).number_format = "0.00%"
        ws.cell(r, 13).number_format = "0.00%"

    for col_idx, hex_color in WARNA_HDR.items():
        fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        for row_idx in range(2, 4):
            ws.cell(row_idx, col_idx).fill = fill

    for col_idx, hex_color in WARNA_DATA.items():
        fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        for row_idx in range(4, total_row + 1):
            ws.cell(row_idx, col_idx).fill = fill

    for row_idx in range(1, total_row + 1):
        for col_idx in range(1, total_col + 1):
            cell = ws.cell(row_idx, col_idx)
            cell.border = Border(
                left   = medium if col_idx == 1 else thin,
                right  = medium if col_idx == total_col else thin,
                top    = medium if row_idx == 1 else thin,
                bottom = medium if (row_idx == total_row or row_idx == 3) else thin
            )
            if row_idx <= 3:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = Font(bold=True)
            elif col_idx == 3:
                # Kolom C (Nama Dealer) rata kiri
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Judul - warna biru tua, font putih tebal
    fill_judul = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    ws["A1"].fill = fill_judul
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    widths = {"A": 12, "B": 8, "C": 26, "D": 12, "E": 16, "F": 11,
              "G": 16, "H": 11, "I": 13, "J": 12, "K": 13, "L": 13, "M": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 32

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def process(file_bytes, bulan, tahun):
    df_raw     = pd.read_excel(io.BytesIO(file_bytes), header=None)
    data       = _parse_header(df_raw)
    data_clean, kolom_md = _clean_data(data)
    profiling  = _aggregate(data_clean, kolom_md)
    excel_bytes= _build_excel(profiling, bulan, tahun)
    filename   = f"Inbound_Website_{bulan}_{tahun}.xlsx"

    preview_rows = []
    for _, row in profiling.iterrows():
        lead     = int(row.get("MD_LEAD", 0))
        wl_d     = int(row.get("MD_WL_D", 0))
        wl_flp   = int(row.get("MD_WL_FLP_Blm_Input_Prospek", 0))
        lapsopir = int(row.get("MD_LAPSOPIR", 0))
        kacab_pct = (wl_d / lead * 100) if lead > 0 else 0
        flp_pct   = (wl_flp / lead * 100) if lead > 0 else 0
        wl_actual = kacab_pct + flp_pct
        cr_actual = (lapsopir / lead * 100) if lead > 0 else 0
        preview_rows.append([
            bulan, tahun, row["Nama_Dealer"], lead,
            wl_d, f"{kacab_pct:.2f}%",
            wl_flp, f"{flp_pct:.2f}%",
            "Maksimal 5%", f"{wl_actual:.2f}%",
            lapsopir, "Minimum 2%", f"{cr_actual:.2f}%"
        ])

    return preview_rows, excel_bytes, filename
