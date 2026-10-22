"""
Processor untuk data Outbound
"""
import io
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


WARNA_HDR = {
    1: "C6EFCE", 2: "DAFDE0", 3: "D9D9D9", 4: "BDD7EE",
    5: "FCE4D6", 6: "FCE4D6", 7: "E2EFDA", 8: "E2EFDA",
    9: "E8D5F5", 10: "E8D5F5", 11: "FFF2CC",
    12: "DDEBF7", 13: "DDEBF7", 14: "FFD7D7",
}
WARNA_DATA = {
    1: "EBF5EB", 2: "F0FAF0", 3: "F2F2F2", 4: "DEEAF1",
    5: "FEF3EE", 6: "FEF3EE", 7: "F2F8EE", 8: "F2F8EE",
    9: "F5EEFA", 10: "F5EEFA", 11: "FFFAE6",
    12: "EEF5FB", 13: "EEF5FB", 14: "FFEEEE",
}

HTML_WARNA_HDR  = {k: f"#{v}" for k, v in WARNA_HDR.items()}
HTML_WARNA_DATA = {k: f"#{v}" for k, v in WARNA_DATA.items()}

# Label kuartal pada sub-header (sesuai notebook asli)
KUARTAL = "Q3"


def _parse_header(df_raw):
    hdr  = df_raw.iloc[1:4].copy()
    data = df_raw.iloc[4:].copy()
    hdr  = hdr.apply(lambda col: col.map(
        lambda x: " ".join(str(x).split()) if pd.notna(x) else np.nan))
    hdr.iloc[0] = hdr.iloc[0].ffill()
    blok  = hdr.iloc[0]
    tahap = hdr.iloc[1].copy()
    for b in blok.dropna().unique():
        idx = blok[blok == b].index
        tahap.loc[idx] = tahap.loc[idx].ffill()
    hdr.iloc[1] = tahap

    def buat_nama(col):
        bagian = [str(x) for x in hdr[col] if pd.notna(x)]
        return "_".join(bagian)

    data.columns = [buat_nama(c) for c in hdr.columns]
    data.columns = [c.replace(" ", "_").replace(".", "") for c in data.columns]
    return data


def _clean_data(data):
    data = data[data.iloc[:, 1].notna()].reset_index(drop=True)

    # Hapus kolom MD & OFFLINE (outbound pakai PROFILLING)
    kolom_md  = [c for c in data.columns if c.startswith("MD")]
    kolom_off = [c for c in data.columns if c.startswith("OFFLINE")]
    data_clean = data.drop(columns=kolom_md + kolom_off)

    # Ambil kolom PROFILLING
    kolom_prof = [c for c in data_clean.columns if c.startswith("PROFILLING")]
    prof = data_clean[kolom_prof].apply(pd.to_numeric, errors="coerce")
    mask_nol   = (prof.fillna(0) == 0).all(axis=1)
    data_clean = data_clean.loc[~mask_nol].reset_index(drop=True)

    for col in kolom_prof:
        data_clean[col] = pd.to_numeric(data_clean[col], errors="coerce").fillna(0)

    return data_clean, kolom_prof


def _aggregate(data_clean, kolom_prof):
    agg_cols = {}
    for col in ["PROFILLING_LEAD", "PROFILLING_WL_D",
                "PROFILLING_WL_FLP_Blm_Input_Prospek", "PROFILLING_LAPSOPIR"]:
        if col in kolom_prof:
            agg_cols[col] = "sum"
    profiling = (
        data_clean.groupby("Nama_Dealer").agg(agg_cols)
        .reset_index()
        .sort_values("PROFILLING_LEAD", ascending=False)
        .reset_index(drop=True)
    )
    return profiling


def _get_sales(df_sales_bytes):
    """Baca file sales, kembalikan dict {DEALER_UPPER: SALES}."""
    df_sales = pd.read_excel(io.BytesIO(df_sales_bytes), header=0)
    cols_upper = {str(c).strip().upper(): c for c in df_sales.columns}
    # Prioritaskan nama kolom persis seperti di notebook: DEALER & SALES
    dealer_col = cols_upper.get("DEALER")
    sales_col  = cols_upper.get("SALES")
    if dealer_col is None:
        dealer_col = next((c for k, c in cols_upper.items() if "DEALER" in k), None)
    if sales_col is None:
        sales_col = next((c for k, c in cols_upper.items() if "SALES" in k), None)
    if dealer_col is None or sales_col is None:
        return {}
    df_sales[sales_col] = pd.to_numeric(df_sales[sales_col], errors="coerce").fillna(0)
    df_sales["_KEY"] = df_sales[dealer_col].astype(str).str.upper().str.strip()
    agg = df_sales.groupby("_KEY")[sales_col].sum()
    return agg.to_dict()


def _build_excel(profiling, bulan, tahun, sales_dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "Profiling Q3"

    # Header baris 1-2 (tanpa baris judul, sesuai notebook)
    ws.merge_cells("A1:A2"); ws.merge_cells("B1:B2")
    ws.merge_cells("C1:C2"); ws.merge_cells("D1:D2")
    ws.merge_cells("E1:F1"); ws.merge_cells("G1:H1")
    ws.merge_cells("I1:J1"); ws.merge_cells("K1:K2")
    ws.merge_cells("L1:M1"); ws.merge_cells("N1:N2")

    ws["A1"] = "Bulan";           ws["B1"] = "Tahun"
    ws["C1"] = "Nama Dealer";     ws["D1"] = "Total Leads"
    ws["E1"] = "KACAB";           ws["G1"] = "FLP"
    ws["I1"] = "Total Workload";  ws["K1"] = "Laporan Sopir"
    ws["L1"] = "Contribution to Sales"; ws["N1"] = "Sales"

    ws["E2"] = "Belum Difollow Up"; ws["F2"] = "Persentase"
    ws["G2"] = "Belum Difollow Up"; ws["H2"] = "Persentase"
    ws["I2"] = f"Target {KUARTAL}";  ws["J2"] = f"Actual {KUARTAL}"
    ws["L2"] = f"Target {KUARTAL}";  ws["M2"] = f"Actual {KUARTAL}"

    for i, row in profiling.iterrows():
        r = i + 3
        lead     = int(row.get("PROFILLING_LEAD", 0))
        wl_d     = int(row.get("PROFILLING_WL_D", 0))
        wl_flp   = int(row.get("PROFILLING_WL_FLP_Blm_Input_Prospek", 0))
        lapsopir = int(row.get("PROFILLING_LAPSOPIR", 0))
        nama_upper = str(row["Nama_Dealer"]).upper().strip()
        sales_val  = int(sales_dict.get(nama_upper, 0))

        ws.cell(r, 1, bulan);  ws.cell(r, 2, tahun)
        ws.cell(r, 3, row["Nama_Dealer"])
        ws.cell(r, 4, lead)
        ws.cell(r, 5, wl_d)
        ws.cell(r, 6, f"=E{r}/D{r}" if lead > 0 else 0)
        ws.cell(r, 7, wl_flp)
        ws.cell(r, 8, f"=G{r}/D{r}" if lead > 0 else 0)
        ws.cell(r, 9, "Maksimal 5%")
        ws.cell(r, 10, f"=F{r}+H{r}")
        ws.cell(r, 11, lapsopir)
        ws.cell(r, 12, "Minimum 10%")
        ws.cell(r, 13, f"=K{r}/N{r}" if sales_val > 0 else 0)
        ws.cell(r, 14, sales_val)

    thin   = Side(style="thin",   color="000000")
    medium = Side(style="medium", color="000000")
    total_col = 14
    total_row = ws.max_row

    for r in range(3, total_row + 1):
        ws.cell(r, 6).number_format  = "0.00%"
        ws.cell(r, 8).number_format  = "0.00%"
        ws.cell(r, 10).number_format = "0.00%"
        ws.cell(r, 13).number_format = "0.00%"

    for col_idx, hex_color in WARNA_HDR.items():
        fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        for row_idx in range(1, 3):
            ws.cell(row_idx, col_idx).fill = fill

    for col_idx, hex_color in WARNA_DATA.items():
        fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        for row_idx in range(3, total_row + 1):
            ws.cell(row_idx, col_idx).fill = fill

    for row_idx in range(1, total_row + 1):
        for col_idx in range(1, total_col + 1):
            cell = ws.cell(row_idx, col_idx)
            cell.border = Border(
                left   = medium if col_idx == 1 else thin,
                right  = medium if col_idx == total_col else thin,
                top    = medium if row_idx == 1 else thin,
                bottom = medium if (row_idx == total_row or row_idx == 2) else thin
            )
            if row_idx <= 2:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.font = Font(bold=True)
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {"A": 12, "B": 8, "C": 22, "D": 12,
              "E": 14, "F": 11, "G": 14, "H": 11,
              "I": 13, "J": 12, "K": 13,
              "L": 15, "M": 12, "N": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 32

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def process(file_bytes, sales_bytes, bulan, tahun):
    df_raw     = pd.read_excel(io.BytesIO(file_bytes), header=None)
    data       = _parse_header(df_raw)
    data_clean, kolom_prof = _clean_data(data)
    profiling  = _aggregate(data_clean, kolom_prof)
    sales_dict = _get_sales(sales_bytes) if sales_bytes else {}
    excel_bytes= _build_excel(profiling, bulan, tahun, sales_dict)
    filename   = f"Outbound_{bulan}_{tahun}.xlsx"

    preview_rows = []
    for _, row in profiling.iterrows():
        lead     = int(row.get("PROFILLING_LEAD", 0))
        wl_d     = int(row.get("PROFILLING_WL_D", 0))
        wl_flp   = int(row.get("PROFILLING_WL_FLP_Blm_Input_Prospek", 0))
        lapsopir = int(row.get("PROFILLING_LAPSOPIR", 0))
        nama_upper = str(row["Nama_Dealer"]).upper().strip()
        sales_val  = int(sales_dict.get(nama_upper, 0))
        kacab_pct = (wl_d / lead * 100) if lead > 0 else 0
        flp_pct   = (wl_flp / lead * 100) if lead > 0 else 0
        wl_actual = kacab_pct + flp_pct
        contrib   = (lapsopir / sales_val * 100) if sales_val > 0 else 0
        preview_rows.append([
            bulan, tahun, row["Nama_Dealer"], lead,
            wl_d, f"{kacab_pct:.2f}%",
            wl_flp, f"{flp_pct:.2f}%",
            "Maksimal 5%", f"{wl_actual:.2f}%",
            lapsopir, "Minimum 10%", f"{contrib:.2f}%",
            sales_val
        ])

    return preview_rows, excel_bytes, filename
